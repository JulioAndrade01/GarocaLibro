import logging
from datetime import datetime, date
from io import BytesIO
import pandas as pd
from openpyxl.utils import get_column_letter
from openpyxl.styles import Font, Alignment, PatternFill

from django.views.generic import TemplateView, CreateView, DeleteView, UpdateView, ListView, DeleteView
from django.urls import reverse_lazy
from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth import authenticate, login
from django.contrib.auth.decorators import login_required, user_passes_test
from django.contrib import messages
from django.http import JsonResponse, HttpResponse
from django.utils.timezone import make_aware
from django.views.decorators.http import require_POST
from django.db.models import Count

from core.forms import LoginForm, LeitorModelForm, AgendamentoForm, LivroModelForm
from core.models import Emprestimo, Leitor, Livro, Agendamento, Noticia

# ===========================================
# 🔹 CONFIGURAÇÃO DE LOG
# ===========================================
logger = logging.getLogger(__name__)

# ===========================================
# 🔹 HOME E APP PRINCIPAL
# ===========================================
def home(request):
    noticias = Noticia.objects.filter(ativo=True).order_by('-criado_em')[:5]
    return render(request, 'home.html', {'noticias': noticias})


class AppGarocaView(TemplateView):
    template_name = 'appgaroca.html'


# ===========================================
# 🔹 CRUD: LEITOR
# ===========================================
class LeitorListView(ListView):
    model = Leitor
    template_name = 'leitor-list.html'
    context_object_name = 'leitores'
    ordering = ['nome']


class LeitorCreateView(CreateView):
    model = Leitor
    form_class = LeitorModelForm
    template_name = 'leitor.html'
    success_url = reverse_lazy('leitor-list')


class LeitorUpdateView(UpdateView):
    model = Leitor
    form_class = LeitorModelForm
    template_name = 'leitor.html'
    success_url = reverse_lazy('leitor-list')



class LeitorDeleteView(DeleteView):
    model = Leitor
    template_name = 'leitor_confirm_delete.html'
    success_url = reverse_lazy('leitor-list')


# ===========================================
# 🔹 CRUD: LIVRO
# ===========================================
class LivroListView(TemplateView):
    template_name = 'livro-list.html'

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['livros'] = Livro.objects.all().order_by('nome')
        return context


class LivroCreateView(CreateView):
    model = Livro
    form_class = LivroModelForm
    template_name = 'livro.html'
    success_url = reverse_lazy('livro-list')


class LivroUpdateView(UpdateView):
    model = Livro
    form_class = LivroModelForm
    template_name = 'livro.html'
    success_url = reverse_lazy('livro-list')


class LivroDeleteView(DeleteView):
    model = Livro
    template_name = 'livro_confirm_delete.html'
    success_url = reverse_lazy('livro-list')


# ===========================================
# 🔹 LISTAGEM SIMPLES DE LIVROS
# ===========================================
def livros_view(request):
    livros = Livro.objects.all().order_by('nome')
    return render(request, 'livros.html', {'livros': livros})


# ===========================================
# 🔹 EMPRÉSTIMO
# ===========================================
class EmprestimoListView(ListView):
    model = Emprestimo
    template_name = 'emprestimo-list.html'
    context_object_name = 'emprestimos'
    ordering = ['-criado']


class EmprestimoCreateView(CreateView):
    model = Emprestimo
    fields = ['livro', 'devolucao']
    template_name = 'emprestimo.html'
    success_url = reverse_lazy('emprestimo-list')

    def form_valid(self, form):
        devolucao = form.cleaned_data['devolucao']
        if isinstance(devolucao, date) and not isinstance(devolucao, datetime):
            form.instance.devolucao = make_aware(datetime.combine(devolucao, datetime.min.time()))
        return super().form_valid(form)

class EmprestimoUpdateView(UpdateView):
    model = Emprestimo
    fields = ['livro', 'devolucao', 'status']
    template_name = 'emprestimo.html'
    success_url = reverse_lazy('emprestimo-list')


# ===========================================
# 🔹 EXCLUSÃO DE EMPRÉSTIMO
# ===========================================w

class EmprestimoDeleteView(DeleteView):
    model = Emprestimo
    template_name = 'emprestimo_confirm_delete.html'
    success_url = reverse_lazy('emprestimo-list')

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['titulo'] = "Excluir Empréstimo"
        context['mensagem'] = f"Deseja realmente excluir o empréstimo do livro '{self.object.livro.nome}'?"# type: ignore
        return context

# ===========================================
# 🔹 AGENDAMENTOS
# ===========================================
@login_required
def reservas_view(request):
    reservas = Emprestimo.objects.filter(leitor=request.user)
    return render(request, 'reservas.html', {'reservas': reservas})


@login_required
def agendar_retirada(request):
    livros_disponiveis = Livro.objects.filter(status=True)

    if request.method == "POST":
        livro_id = request.POST.get("livro")
        data_agendada = request.POST.get("data_agendada")

        if not livro_id or not data_agendada:
            messages.error(request, "Por favor, selecione um livro e uma data.")
            return redirect("agendar_retirada")

        livro = get_object_or_404(Livro, id=livro_id, status=True)
        Agendamento.objects.create(
            leitor=request.user,
            livro=livro,
            data_agendada=data_agendada,
            status="scheduled",
        )

        livro.status = False
        livro.save(update_fields=["status"])

        messages.success(
            request,
            f"📚 Agendamento realizado com sucesso para o livro '{livro.nome}' em {data_agendada}!"
        )
        return redirect("dashboard_leitor")

    context = {
        "livros_list": livros_disponiveis,
        "today": date.today().isoformat(),
    }
    return render(request, "agendar_retirada.html", context)


@login_required
@require_POST
def cancelar_agendamento(request, agendamento_id):
    try:
        agendamento = get_object_or_404(Agendamento, id=agendamento_id, leitor=request.user)
        if agendamento.status != 'scheduled':
            return JsonResponse({'error': 'Este agendamento não pode ser cancelado.'}, status=400)

        agendamento.status = 'cancelled'
        agendamento.save()

        agendamento.livro.status = True
        agendamento.livro.save()

        return JsonResponse({'success': True, 'message': 'Agendamento cancelado com sucesso!'})
    except Exception as e:
        logger.error(f"Erro ao cancelar agendamento: {str(e)}")
        return JsonResponse({'error': 'Erro ao cancelar o agendamento.'}, status=500)


# ===========================================
# 🔹 LOGIN / CADASTRO / PERFIL
# ===========================================
def register(request):
    if request.method == 'POST':
        form = LeitorModelForm(request.POST)
        if form.is_valid():
            leitor = form.save(commit=False)
            leitor.telefone = 'Desconhecido'
            leitor.set_password(form.cleaned_data['password'])
            leitor.save()
            login(request, leitor)
            return redirect('home')
    else:
        form = LeitorModelForm()
    return render(request, 'register.html', {'form': form})


def login_view(request):
    form = LoginForm(request.POST or None)
    if request.method == "POST" and form.is_valid():
        email = form.cleaned_data['email']
        password = form.cleaned_data['password']
        user = authenticate(request, email=email, password=password)
        if user is not None:
            login(request, user)
            return redirect('dashboard_leitor')
        else:
            messages.error(request, "Credenciais inválidas.")

    response = render(request, 'login.html', {'form': form})
    response['Cache-Control'] = 'no-store, no-cache, must-revalidate, proxy-revalidate'
    return response


@login_required
def perfil_view(request):
    leitor = get_object_or_404(Leitor, email=request.user.email)
    return render(request, 'perfil.html', {'leitor': leitor})


@login_required
def editar_perfil_view(request):
    leitor = get_object_or_404(Leitor, email=request.user.email)
    if request.method == 'POST':
        form = LeitorModelForm(request.POST, request.FILES, instance=leitor)
        if form.is_valid():
            form.save()
            messages.success(request, 'Perfil atualizado com sucesso!')
            return redirect('perfil')
    else:
        form = LeitorModelForm(instance=leitor)
    return render(request, 'editar_perfil.html', {'form': form})


@login_required
def editar_perfil_ajax(request):
    leitor = request.user
    if request.method == "POST":
        nome = request.POST.get("nome")
        telefone = request.POST.get("telefone")
        foto = request.FILES.get("foto_perfil")

        if nome:
            leitor.nome = nome
        if telefone:
            leitor.telefone = telefone
        if foto:
            leitor.foto_perfil = foto

        leitor.save()

        return JsonResponse({
            "success": True,
            "foto_perfil": leitor.foto_perfil.url if leitor.foto_perfil else "/static/img/default-profile.png"
        })

    return JsonResponse({"success": False, "error": "Requisição inválida."}, status=400)


# ===========================================
# 🔹 DASHBOARD LEITOR
# ===========================================
@login_required
def dashboard_leitor(request):
    leitor = request.user
    emprestimos = Emprestimo.objects.filter(leitor=leitor).order_by('-criado')
    agendamentos = Agendamento.objects.filter(leitor=leitor).order_by('-criado')

    context = {
        'leitor': leitor,
        'emprestimos': emprestimos,
        'agendamentos': agendamentos,
    }
    return render(request, 'leitor_dashboard/dashboard.html', context)


# ===========================================
# 🔹 MEUS EMPRÉSTIMOS
# ===========================================
@login_required
def meus_emprestimos(request):
    emprestimos = Emprestimo.objects.filter(leitor=request.user).order_by('-criado')
    return render(request, 'leitor_dashboard/meus_emprestimos.html', {'emprestimos': emprestimos})


# ===========================================
# 🔹 DEVOLUÇÃO DE LIVROS
# ===========================================
@login_required
def devolver_livro_view(request):
    if request.method == 'POST':
        codigo_livro = request.POST.get('codigo_livro')
        try:
            livro = get_object_or_404(Livro, codigo=codigo_livro)
            emprestimo = get_object_or_404(Emprestimo, livro=livro, status='in_progress')
            emprestimo.status = 'completed'
            emprestimo.save()
            livro.status = True
            livro.save()
            messages.success(request, f'📚 Livro "{livro.nome}" devolvido com sucesso!')
            return redirect('devolver_livro')
        except Livro.DoesNotExist:
            messages.error(request, 'Código do livro não encontrado.')
        except Emprestimo.DoesNotExist:
            messages.error(request, 'Este livro não está emprestado.')
        except Exception as e:
            logger.error(f"Erro ao processar devolução: {str(e)}")
            messages.error(request, 'Erro ao processar devolução.')
    return render(request, 'devolver_livro.html')


@login_required
def api_devolver_livro(request, emprestimo_id):
    try:
        emprestimo = get_object_or_404(Emprestimo, id=emprestimo_id, leitor=request.user)
        if emprestimo.status != 'in_progress':
            return JsonResponse({'error': 'Este empréstimo já foi finalizado.'}, status=400)

        emprestimo.status = 'completed'
        emprestimo.save()

        livro = emprestimo.livro
        livro.status = True
        livro.save()

        return JsonResponse({'message': 'Livro devolvido com sucesso!'})
    except Exception as e:
        logger.error(f"Erro ao processar devolução via API: {str(e)}")
        return JsonResponse({'error': 'Erro ao processar devolução.'}, status=500)


# ===========================================
# 📊 ANALYTICS (ADMIN)
# ===========================================
@user_passes_test(lambda u: u.is_staff)
def analytics_dashboard(request):
    top_livros = Emprestimo.objects.values('livro__nome').annotate(total=Count('livro')).order_by('-total')[:10]
    top_bairros = Leitor.objects.values('bairro').annotate(total=Count('id')).order_by('-total')[:10]
    emprestimos_mensais = (
        Emprestimo.objects.extra(select={'month': "strftime('%%m', criado)"})
        .values('month').annotate(total=Count('id')).order_by('month')
    )

    context = {
        'top_livros': list(top_livros),
        'top_bairros': list(top_bairros),
        'emprestimos_mensais': list(emprestimos_mensais),
        'total_livros': Livro.objects.count(),
        'total_leitores': Leitor.objects.count(),
        'total_emprestimos': Emprestimo.objects.count(),
        'total_bairros': Leitor.objects.values('bairro').distinct().count(),
    }
    return render(request, 'analytics/dashboard_admin.html', context)


@user_passes_test(lambda u: u.is_staff)
def analytics_report(request):
    bairros = Leitor.objects.values('bairro').exclude(bairro__isnull=True).distinct().order_by('bairro')
    livros = Livro.objects.values('nome').order_by('nome')
    return render(request, 'analytics/report.html', {'bairros': bairros, 'livros': livros})


@user_passes_test(lambda u: u.is_staff)
def filtrar_emprestimos(request):
    data_inicio = request.GET.get('data_inicio')
    data_fim = request.GET.get('data_fim')
    bairro = request.GET.get('bairro')
    livro = request.GET.get('livro')

    emprestimos = Emprestimo.objects.select_related('livro', 'leitor')

    if data_inicio:
        emprestimos = emprestimos.filter(criado__gte=data_inicio)
    if data_fim:
        emprestimos = emprestimos.filter(criado__lte=data_fim)
    if bairro:
        emprestimos = emprestimos.filter(leitor__bairro=bairro)
    if livro:
        emprestimos = emprestimos.filter(livro__nome=livro)

    resultados = [
        {
            'livro': e.livro.nome,
            'leitor': e.leitor.nome,
            'bairro': e.leitor.bairro,
            'data_emprestimo': e.criado.strftime('%d/%m/%Y'),
            'data_devolucao': e.devolucao.strftime('%d/%m/%Y') if e.devolucao else None,
            'status': e.status,
        }
        for e in emprestimos
    ]

    return JsonResponse({'resultados': resultados})


@user_passes_test(lambda u: u.is_staff)
def insights_dashboard(request):
    top_livros = Emprestimo.objects.values('livro__nome').annotate(total=Count('livro')).order_by('-total')[:5]
    top_bairros = Leitor.objects.values('bairro').annotate(total=Count('id')).order_by('-total')[:5]
    emprestimos_mensais = (
        Emprestimo.objects.extra(select={'month': "strftime('%%m', criado)"})
        .values('month').annotate(total=Count('id')).order_by('month')
    )

    context = {
        'total_livros': Livro.objects.count(),
        'total_leitores': Leitor.objects.count(),
        'total_emprestimos': Emprestimo.objects.count(),
        'total_bairros': Leitor.objects.values('bairro').distinct().count(),
        'top_livros': list(top_livros),
        'top_bairros': list(top_bairros),
        'emprestimos_mensais': list(emprestimos_mensais),
    }
    return render(request, 'analytics/insights.html', context)


@user_passes_test(lambda u: u.is_staff) # type: ignore
def exportar_relatorio_excel(request):
    data_inicio = request.GET.get('data_inicio')
    data_fim = request.GET.get('data_fim')
    bairro = request.GET.get('bairro')
    livro_nome = request.GET.get('livro')

    emprestimos = Emprestimo.objects.select_related('livro', 'leitor')

    if data_inicio:
        emprestimos = emprestimos.filter(criado__gte=data_inicio)
    if data_fim:
        emprestimos = emprestimos.filter(criado__lte=data_fim)
    if bairro:
        emprestimos = emprestimos.filter(leitor__bairro=bairro)
    if livro_nome:
        emprestimos = emprestimos.filter(livro__nome=livro_nome)

    if not emprestimos.exists():
        messages.warning(request, "Nenhum registro encontrado com esses filtros.")
        return redirect('analytics_report')

    df = pd.DataFrame(list(emprestimos.values(
        'livro__nome',
        'leitor__nome',
        'leitor__bairro',
        'criado',
        'devolucao',
        'status'
    )))

    df.rename(columns={
        'livro__nome': 'Livro',
        'leitor__nome': 'Leitor',
        'leitor__bairro': 'Bairro',
        'criado': 'Data Empréstimo',
        'devolucao': 'Data Devolução',
        'status': 'Status'
    }, inplace=True)

    for coluna in ['Data Empréstimo', 'Data Devolução']:
        df[coluna] = pd.to_datetime(df[coluna], errors='coerce')

    df['Data Empréstimo'] = df['Data Empréstimo'].dt.strftime('%d/%m/%Y')
    df['Data Devolução'] = df['Data Devolução'].dt.strftime('%d/%m/%Y')

    output = BytesIO()
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        df.to_excel(writer, index=False, sheet_name='Empréstimos')
        ws = writer.sheets['Empréstimos']

        header_font = Font(bold=True, color="FFFFFF")
        fill = PatternFill(start_color="0070C0", end_color="0070C0", fill_type="solid")

        for col_num, column_title in enumerate(df.columns, 1):
            cell = ws.cell(row=1, column=col_num)
            cell.font = header_font
            cell.fill = fill
            cell.alignment = Alignment(horizontal="center")
            ws.column_dimensions[get_column_letter(col_num)].width = max(len(column_title) + 5, 18)

